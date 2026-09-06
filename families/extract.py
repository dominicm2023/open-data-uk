"""Bounded extraction of one fetched file into tables. Subprocess; no network.

Derived from Codex's pilot extractor (September 2026), with two formats
added: an ArcGIS REST query result and GeoJSON, both of which arrive as
features rather than rows and are flattened here so every family adapter
sees the same shape — a header row followed by value rows. Geometry becomes
`_lon`/`_lat` columns (WGS84, because the query asks for outSR=4326) so an
adapter can treat coordinates like any other column.

Usage: extract.py <blob> <FORMAT> <out.json> '<limits json>'
"""
from __future__ import annotations

import csv
import io
import json
import sys
import zipfile
from pathlib import Path

VERSION = "tables-v2"


def _features_to_rows(features: list[dict], props_key: str, limits: dict) -> list[list]:
    keys: list[str] = []
    seen = set()
    for f in features:
        for k in (f.get(props_key) or {}):
            if k not in seen:
                seen.add(k)
                keys.append(k)
    header = keys + ["_lon", "_lat"]
    rows = [header]
    for f in features:
        if len(rows) > limits["max_rows"]:
            raise ValueError("Row limit exceeded")
        p = f.get(props_key) or {}
        lon = lat = None
        g = f.get("geometry") or {}
        if "x" in g and "y" in g:                       # ArcGIS point
            lon, lat = g.get("x"), g.get("y")
        elif g.get("type") == "Point" and g.get("coordinates"):
            lon, lat = g["coordinates"][:2]
        rows.append([p.get(k) for k in keys] + [lon, lat])
    return rows


def extract(path: Path, fmt: str, limits: dict) -> list[dict]:
    tables: list[dict] = []
    total = 0

    def add(rows, **where):
        nonlocal total
        total += len(rows)
        if total > limits["max_rows"]:
            raise ValueError("Row limit exceeded")
        tables.append(dict(**where, rows=rows))

    data = path.read_bytes()
    if fmt == "CSV":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp1252")
        if text.lstrip().lower().startswith(("<!doctype", "<html")):
            raise ValueError("HTML response, not CSV")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = []
        for row in csv.reader(io.StringIO(text), dialect):
            if len(rows) >= limits["max_rows"]:
                raise ValueError("Row limit exceeded")
            rows.append(row)
        if not rows or max(map(len, rows)) < 2:
            raise ValueError("No tabular CSV content")
        add(rows, table=1)
    elif fmt == "ESRI":
        doc = json.loads(data)
        if doc.get("error"):
            raise ValueError("ArcGIS error: " + str(doc["error"])[:200])
        feats = doc.get("features")
        if not isinstance(feats, list):
            raise ValueError("No features in ArcGIS response")
        add(_features_to_rows(feats, "attributes", limits), table=1,
            exceeded=bool(doc.get("exceededTransferLimit")))
    elif fmt == "GEOJSON":
        doc = json.loads(data)
        feats = doc.get("features")
        if not isinstance(feats, list):
            raise ValueError("Not a GeoJSON FeatureCollection")
        add(_features_to_rows(feats, "properties", limits), table=1)
    elif fmt == "JSON":
        doc = json.loads(data)
        if isinstance(doc, dict) and doc.get("error"):
            raise ValueError("API returned an error")
        if isinstance(doc, dict) and isinstance(doc.get("features"), list):
            key = "attributes" if doc["features"] and "attributes" in doc["features"][0] else "properties"
            add(_features_to_rows(doc["features"], key, limits), table=1)
        elif isinstance(doc, list) and doc and isinstance(doc[0], dict):
            keys = list(dict.fromkeys(k for r in doc for k in r))
            add([keys] + [[r.get(k) for k in keys] for r in doc], table=1)
        else:
            raise ValueError("JSON is not a list of records or a feature collection")
    elif fmt == "XLSX":
        import openpyxl
        with zipfile.ZipFile(path) as archive:
            if sum(z.file_size for z in archive.infolist()) > 100_000_000:
                raise ValueError("Expanded XLSX exceeds 100 MB")
        with path.open("rb") as handle:
            book = openpyxl.load_workbook(handle, read_only=True, data_only=True)
            for sheet in book:
                if sheet.max_column and sheet.max_column > 100:
                    raise ValueError("Column limit exceeded")
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    if len(rows) + total >= limits["max_rows"]:
                        raise ValueError("Row limit exceeded")
                    rows.append([("" if v is None else v) for v in row])
                if rows:
                    add(rows, sheet=sheet.title)
            book.close()
    elif fmt == "PDF":
        import pdfplumber
        if data[:5] != b"%PDF-":
            raise ValueError("Not a PDF")
        with pdfplumber.open(path) as pdf:
            if len(pdf.pages) > limits["max_pages"]:
                raise ValueError("Page limit exceeded")
            for pno, page in enumerate(pdf.pages, 1):
                for tno, table in enumerate(page.find_tables(), 1):
                    add(table.extract(), page=pno, table=tno, bbox=list(table.bbox))
        if not tables:
            raise ValueError("No ruled tables; layout adapter required")
    else:
        raise ValueError(f"Unsupported format {fmt}")
    return tables


def main() -> int:
    blob, fmt, out, limits = sys.argv[1], sys.argv[2], sys.argv[3], json.loads(sys.argv[4])
    tables = extract(Path(blob), fmt, limits)
    payload = json.dumps({"version": VERSION, "format": fmt, "tables": tables},
                         ensure_ascii=False, default=str)
    if len(payload) > limits["max_output_bytes"]:
        raise ValueError("Extraction output exceeds byte limit")
    Path(out).write_text(payload, encoding="utf-8")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:  # noqa: BLE001 - the parent reads stderr
        print(str(exc), file=sys.stderr)
        raise SystemExit(1)
